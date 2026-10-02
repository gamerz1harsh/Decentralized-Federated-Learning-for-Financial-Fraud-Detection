"""
Generate a concise PPT (max 7 slides) for the project:
Abstract, Introduction, and figure-based explanations.
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
import os

FIGURES_DIR = os.path.join(os.path.dirname(__file__), "figures")
OUTPUT_PATH = os.path.join(os.path.dirname(__file__), "Project_Presentation.pptx")

# Colors
DARK_BLUE = RGBColor(0x1F, 0x3A, 0x5F)
MID_BLUE = RGBColor(0x2E, 0x5C, 0x8A)
LIGHT_BLUE = RGBColor(0xDE, 0xEB, 0xF7)
ACCENT = RGBColor(0xC5, 0x5A, 0x11)
GRAY = RGBColor(0x44, 0x44, 0x44)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

BLANK = prs.slide_layouts[6]


def add_bg(slide, color):
    """Set slide background color."""
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_textbox(slide, left, top, width, height, text, font_size=18,
                bold=False, color=DARK_BLUE, align=PP_ALIGN.LEFT, font_name="Calibri"):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.font.name = font_name
    p.alignment = align
    return txBox


def add_bullets(slide, left, top, width, height, items, font_size=16, color=DARK_BLUE):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = f"\u2022  {item}"
        p.font.size = Pt(font_size)
        p.font.color.rgb = color
        p.font.name = "Calibri"
        p.space_after = Pt(6)
    return txBox


def add_figure(slide, path, left, top, width=None, height=None):
    if width and height:
        slide.shapes.add_picture(path, left, top, width=width, height=height)
    elif width:
        slide.shapes.add_picture(path, left, top, width=width)
    else:
        slide.shapes.add_picture(path, left, top, height=height)


def add_title_bar(slide, title_text):
    """Add a colored title bar at the top."""
    bar = slide.shapes.add_shape(1, Inches(0), Inches(0), prs.slide_width, Inches(1.0))
    bar.fill.solid()
    bar.fill.fore_color.rgb = DARK_BLUE
    bar.line.fill.background()
    tf = bar.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.font.name = "Calibri"
    p.alignment = PP_ALIGN.LEFT
    tf.margin_left = Inches(0.5)
    tf.margin_top = Inches(0.15)


# ==========================================================
# SLIDE 1: Title
# ==========================================================

slide = prs.slides.add_slide(BLANK)
add_bg(slide, WHITE)

# Decorative top band
band = slide.shapes.add_shape(1, Inches(0), Inches(0), prs.slide_width, Inches(2.2))
band.fill.solid()
band.fill.fore_color.rgb = DARK_BLUE
band.line.fill.background()

add_textbox(slide, Inches(0.8), Inches(0.6), Inches(11.7), Inches(1.2),
            "Federated Learning for Credit Card Fraud Detection",
            font_size=36, bold=True, color=WHITE, align=PP_ALIGN.CENTER)

add_textbox(slide, Inches(0.8), Inches(1.6), Inches(11.7), Inches(0.6),
            "with Multi-Dimensional Client Contribution Scoring",
            font_size=22, bold=False, color=LIGHT_BLUE, align=PP_ALIGN.CENTER)

add_textbox(slide, Inches(0.8), Inches(3.2), Inches(11.7), Inches(0.6),
            "Project Presentation",
            font_size=24, bold=True, color=MID_BLUE, align=PP_ALIGN.CENTER)

add_textbox(slide, Inches(0.8), Inches(4.0), Inches(11.7), Inches(0.5),
            "Abstract  \u2022  Introduction  \u2022  Proposed Approach",
            font_size=18, color=GRAY, align=PP_ALIGN.CENTER)

# ==========================================================
# SLIDE 2: Abstract
# ==========================================================

slide = prs.slides.add_slide(BLANK)
add_bg(slide, WHITE)
add_title_bar(slide, "Abstract")

add_textbox(slide, Inches(0.7), Inches(1.3), Inches(12.0), Inches(5.8),
            "Credit card fraud remains a persistent and financially damaging problem. "
            "We use a local federated-learning simulation to study how multiple clients can "
            "train on separate transaction shards and combine model updates. The simulator "
            "does not implement secure aggregation or differential privacy.",
            font_size=18, color=DARK_BLUE)

add_textbox(slide, Inches(0.7), Inches(2.6), Inches(12.0), Inches(4.5),
            "Clients differ in data size and fraud patterns. We compare FedAvg and robust "
            "baselines with an experimental contribution-aware scorer using validation "
            "quality, update-scale reliability, hard-fraud utility, complementarity, and "
            "history. Results vary by scenario, with no consistent overall gain established.",
            font_size=18, color=DARK_BLUE)

# ==========================================================
# SLIDE 3: Introduction — Background & Problem
# ==========================================================

slide = prs.slides.add_slide(BLANK)
add_bg(slide, WHITE)
add_title_bar(slide, "Introduction: Background & Problem")

add_bullets(slide, Inches(0.7), Inches(1.4), Inches(12.0), Inches(5.8), [
    "Fraud is extremely rare (~0.17% of transactions) and adversarial \u2014 fraudsters constantly adapt.",
    "Centralized fraud detection requires pooling sensitive data, raising privacy and regulatory concerns.",
    "Our simulator keeps client training shards separate and exchanges model updates plus validation metrics.",
    "Problem: In a federation, banks differ in data size, quality, fraud patterns, and reliability.",
    "Standard FedAvg weights all clients equally (or by sample count), ignoring these differences.",
    "Need: A principled way to score each client's contribution and use it to guide aggregation.",
], font_size=17)

# ==========================================================
# SLIDE 4: Proposed Approach — System Architecture (Figure 1)
# ==========================================================

slide = prs.slides.add_slide(BLANK)
add_bg(slide, WHITE)
add_title_bar(slide, "Implemented System: Same-Schema Simulation")

add_figure(slide, os.path.join(FIGURES_DIR, "fig1_system_architecture.png"),
           Inches(0.7), Inches(1.3), width=Inches(7.5))

add_textbox(slide, Inches(8.5), Inches(1.5), Inches(4.3), Inches(5.5),
            "Each client trains from its assigned shard.\n\n"
            "Clients return model updates and validation metrics.\n\n"
            "The server applies the selected aggregation method.\n\n"
            "This simulator does not provide secure aggregation or differential privacy.",
            font_size=16, color=DARK_BLUE)

# ==========================================================
# SLIDE 5: Multi-Dimensional Scoring (Figure 3)
# ==========================================================

slide = prs.slides.add_slide(BLANK)
add_bg(slide, WHITE)
add_title_bar(slide, "Multi-Dimensional Client Scoring")

add_figure(slide, os.path.join(FIGURES_DIR, "fig3_scoring_pipeline.png"),
           Inches(0.7), Inches(1.3), width=Inches(7.5))

add_textbox(slide, Inches(8.5), Inches(1.5), Inches(4.3), Inches(5.5),
            "Quality \u2014 Local validation PR-AUC, adjusted for small fraud samples.\n\n"
            "Reliability \u2014 Robust check of update scale and client history.\n\n"
            "Hard-fraud utility \u2014 Leave-one-out gain on difficult positives.\n\n"
            "Complementarity \u2014 Leave-one-out validation PR-AUC gain.\n\n"
            "Temporal \u2014 Smoothed history of marginal utility.",
            font_size=16, color=DARK_BLUE)

# ==========================================================
# SLIDE 6: Training Workflow (Figure 2)
# ==========================================================

slide = prs.slides.add_slide(BLANK)
add_bg(slide, WHITE)
add_title_bar(slide, "Federated Training Workflow")

add_figure(slide, os.path.join(FIGURES_DIR, "fig2_training_workflow.png"),
           Inches(0.7), Inches(1.3), width=Inches(7.5))

add_textbox(slide, Inches(8.5), Inches(1.5), Inches(4.3), Inches(5.5),
            "1. Server distributes the global model.\n\n"
            "2. Clients train on separate local shards.\n\n"
            "3. Clients return updates and validation metrics.\n\n"
            "4. Server applies the selected aggregation method.\n\n"
            "5. Validation metrics are recorded for the round.\n\n"
            "6. Repeat for the configured rounds.",
            font_size=16, color=DARK_BLUE)

# ==========================================================
# SLIDE 7: Summary & Next Steps
# ==========================================================

slide = prs.slides.add_slide(BLANK)
add_bg(slide, WHITE)
add_title_bar(slide, "Summary & Next Steps")

add_bullets(slide, Inches(0.7), Inches(1.4), Inches(12.0), Inches(5.8), [
    "We implemented a same-schema federated simulator with contribution-aware scoring.",
    "The scorer uses local quality, update reliability, hard-fraud utility, complementarity, and history.",
    "Our local simulator has no secure aggregation or differential privacy.",
    "Results vary by scenario and do not show consistent gains over FedAvg.",
    "Next: expand paired seeds and evaluate on an untouched test set.",
    "Integrate heterogeneous encoders after resolving latent-space alignment.",
], font_size=17)

# ==========================================================
# SAVE
# ==========================================================

prs.save(OUTPUT_PATH)
print(f"Presentation saved to: {OUTPUT_PATH}")
print(f"Number of slides: {len(prs.slides)}")